import sys
sys.path.insert(0, '/Users/javiermartinamador/Documents/Agencia Marketing')

from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

BASE_DIR = '/Users/javiermartinamador/Documents/Agencia Marketing'
BG_DIR = f'{BASE_DIR}/assets/pdf_backgrounds'

prs = Presentation()
prs.slide_width = Cm(29.7)
prs.slide_height = Cm(21.0)

blank_layout = prs.slide_layouts[6]

def add_text_box(slide, left, top, width, height, text, font_size=12, bold=False, color=RGBColor(0xFF,0xFF,0xFF), align=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox

# Slide 1: Cover (dark background)
slide1 = prs.slides.add_slide(blank_layout)
slide1.shapes.add_picture(f'{BG_DIR}/josevicente_slide_1.png', Cm(0), Cm(0), width=Cm(29.7), height=Cm(21.0))

# Editable text overlays (white text on dark bg)
add_text_box(slide1, Cm(2.5), Cm(6.0), Cm(20), Cm(1.2), "JOSE VICENTE SL (B28927994)", 18, True, RGBColor(0xFF,0xFF,0xFF), PP_ALIGN.LEFT)

# Slide 2: Detail table
slide2 = prs.slides.add_slide(blank_layout)
slide2.shapes.add_picture(f'{BG_DIR}/josevicente_slide_2.png', Cm(0), Cm(0), width=Cm(29.7), height=Cm(21.0))

# Header info
add_text_box(slide2, Cm(6.5), Cm(1.0), Cm(12), Cm(0.5), "JOSE VICENTE SL (B28927994)", 11, True, RGBColor(0x00,0x00,0x00), PP_ALIGN.LEFT)
add_text_box(slide2, Cm(6.5), Cm(1.5), Cm(12), Cm(0.3), "Comercial: JAVIER MARTIN AMADOR", 8, False, RGBColor(0x00,0x00,0x00), PP_ALIGN.LEFT)
add_text_box(slide2, Cm(6.5), Cm(1.9), Cm(12), Cm(0.3), "Teléfono: 664258977", 8, False, RGBColor(0x00,0x00,0x00), PP_ALIGN.LEFT)
add_text_box(slide2, Cm(6.5), Cm(2.3), Cm(12), Cm(0.3), "C. Electrónico: JAVIER.MARTIN@LYNKS-TIC.COM", 8, False, RGBColor(0x00,0x00,0x00), PP_ALIGN.LEFT)

# Table data (positioned over the table area)
table_y_start = 5.0
row_height = 0.65

table_rows = [
    ("PAQUETES", "", "B2ONE B8 PREMIUM", "36 meses", "1", "149,00€/mes", "", "", "", "149,00 €"),
    ("", "", "LYNKS FTTO 1GB + BACKUP 4G", "36 meses", "1", "", "", "", "", "0,00 €"),
    ("", "", "CANAL VOZ (TRUNK)", "36 meses", "8", "", "", "", "", "0,00 €"),
    ("", "", "EXT IP (IVR,GRABACIÓN,SOFTPHONE)", "36 meses", "8", "", "", "", "", "0,00 €"),
    ("", "", "DDI (NACIONAL)", "36 meses", "8", "", "", "", "", "0,00 €"),
    ("", "", "MOVIL B2 ILIMITADA PACK", "36 meses", "2", "", "", "", "", "0,00 €"),
    ("LYNKS MOBILE", "TARIFAS", "MOVIL B2 ILIMITADA PACK", "12 meses", "4", "9,45€/mes", "", "", "", "37,80 €"),
    ("LYNKS IP", "EQUIPOS Y TERMINALES", "GIGASET BÁSICO P710", "36 meses", "8", "3,00€/mes", "", "", "", "24,00 €"),
]

col_x = [0.8, 3.3, 5.8, 12.8, 16.3, 18.8, 21.3, 23.8, 26.3, 28.3]

for row_idx, row_data in enumerate(table_rows):
    y = Cm(table_y_start + row_idx * row_height)
    for col_idx, cell_text in enumerate(row_data):
        if cell_text:
            x = Cm(col_x[col_idx])
            w = Cm(2.2) if col_idx < 9 else Cm(1.8)
            align = PP_ALIGN.LEFT if col_idx in [0,1,2] else PP_ALIGN.CENTER
            add_text_box(slide2, x, y, w, Cm(0.5), cell_text, 7, False, RGBColor(0x00,0x00,0x00), align)

# Total row
add_text_box(slide2, Cm(20.5), Cm(10.5), Cm(4), Cm(0.6), "210,80 €", 10, True, RGBColor(0xFF,0xFF,0xFF), PP_ALIGN.RIGHT)
add_text_box(slide2, Cm(24.5), Cm(10.5), Cm(4), Cm(0.6), "210,80 €/mes", 9, True, RGBColor(0xFF,0xFF,0xFF), PP_ALIGN.RIGHT)

# Slide 3: Cloud/Backup/Ciberseguridad
slide3 = prs.slides.add_slide(blank_layout)
slide3.shapes.add_picture(f'{BG_DIR}/josevicente_slide_7.png', Cm(0), Cm(0), width=Cm(29.7), height=Cm(21.0))

# Footer on all slides
for slide in prs.slides:
    add_text_box(slide, Cm(1), Cm(19.5), Cm(25), Cm(0.5), "lynks-tic.com  ·  911 08 98 77  ·  info@lynks-tic.com", 8, False, RGBColor(0xFF,0xFF,0xFF), PP_ALIGN.LEFT)

output_path = f"{BASE_DIR}/templates/lynks-tic_pro_v2_base.pptx"
prs.save(output_path)
print(f"Template saved: {output_path}")
print(f"Slides: {len(prs.slides)}")
print(f"Size: {os.path.getsize(output_path)} bytes")