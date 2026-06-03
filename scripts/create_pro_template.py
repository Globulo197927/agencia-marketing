import sys
sys.path.insert(0, '/Users/javiermartinamador/Documents/Agencia Marketing')

from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

# Design System
DARK_BG = RGBColor(0x1E, 0x1E, 0x1E)
GREEN = RGBColor(0x6A, 0xBD, 0x45)
DARK_GREEN = RGBColor(0x4A, 0x8A, 0x2F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xB0, 0xB0)
MID_GRAY = RGBColor(0x80, 0x80, 0x80)
TABLE_HEADER_BG = RGBColor(0x6A, 0xBD, 0x45)
TABLE_ROW_BG = RGBColor(0xE8, 0xF5, 0xE0)
TABLE_ALT_BG = RGBColor(0xD0, 0xEB, 0xC5)
TOTAL_BG = RGBColor(0x4A, 0x8A, 0x2F)
BLACK = RGBColor(0x00, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Cm(29.7)
prs.slide_height = Cm(21.0)

blank_layout = prs.slide_layouts[6]

def add_text_box(slide, left, top, width, height, text, font_size=12, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font_name='Calibri', bg_color=None):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    if bg_color:
        txBox.fill.solid()
        txBox.fill.fore_color.rgb = bg_color
    return txBox

def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ============================
# SLIDE 1: COVER
# ============================
slide1 = prs.slides.add_slide(blank_layout)

# Dark background
bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(29.7), Cm(21.0))
bg.fill.solid()
bg.fill.fore_color.rgb = DARK_BG
bg.line.fill.background()

# "Propuesta comercial" title
add_text_box(slide1, Cm(2), Cm(3.5), Cm(18), Cm(3), "Propuesta\ncomercial", 54, False, WHITE, PP_ALIGN.LEFT)

# Client name
add_text_box(slide1, Cm(2), Cm(7.5), Cm(20), Cm(1), "JOSE VICENTE SL (B28927994)", 20, True, WHITE, PP_ALIGN.LEFT)

# "lynks" big logo text (right side)
add_text_box(slide1, Cm(16), Cm(11), Cm(12), Cm(2.5), "lynks", 96, False, GREEN, PP_ALIGN.RIGHT)

# "tic" small next to lynks
add_text_box(slide1, Cm(26.5), Cm(12.8), Cm(2), Cm(0.8), "tic", 32, False, GREEN, PP_ALIGN.LEFT)

# "professional network"
add_text_box(slide1, Cm(16), Cm(14.5), Cm(12), Cm(0.6), "professional network", 14, False, LIGHT_GRAY, PP_ALIGN.RIGHT)

# Footer
add_text_box(slide1, Cm(2), Cm(19.5), Cm(10), Cm(0.5), "lynks-tic.com", 10, False, LIGHT_GRAY, PP_ALIGN.LEFT)

# ============================
# SLIDE 2: DETAIL TABLE
# ============================
slide2 = prs.slides.add_slide(blank_layout)

# White background
bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(29.7), Cm(21.0))
bg2.fill.solid()
bg2.fill.fore_color.rgb = WHITE
bg2.line.fill.background()

# Top left logo area (dark box with logo)
logo_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.5), Cm(0.5), Cm(5), Cm(2.5))
logo_box.fill.solid()
logo_box.fill.fore_color.rgb = DARK_BG
logo_box.line.fill.background()

add_text_box(slide2, Cm(0.7), Cm(0.8), Cm(4.6), Cm(0.8), "lynks", 24, False, GREEN, PP_ALIGN.LEFT)
add_text_box(slide2, Cm(3.5), Cm(1.1), Cm(1.5), Cm(0.4), "tic", 10, False, GREEN, PP_ALIGN.LEFT)
add_text_box(slide2, Cm(0.7), Cm(1.6), Cm(4.6), Cm(0.3), "professional network", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT)

# Client name and commercial info
add_text_box(slide2, Cm(6), Cm(0.8), Cm(15), Cm(0.6), "JOSE VICENTE SL (B28927994)", 12, True, BLACK, PP_ALIGN.LEFT)
add_text_box(slide2, Cm(6), Cm(1.4), Cm(15), Cm(0.4), "Comercial: JAVIER MARTIN AMADOR", 9, False, MID_GRAY, PP_ALIGN.LEFT)
add_text_box(slide2, Cm(6), Cm(1.8), Cm(15), Cm(0.4), "Teléfono: 664258977", 9, False, MID_GRAY, PP_ALIGN.LEFT)
add_text_box(slide2, Cm(6), Cm(2.2), Cm(15), Cm(0.4), "C. Electrónico: JAVIER.MARTIN@LYNKS-TIC.COM", 9, False, MID_GRAY, PP_ALIGN.LEFT)

# "DETALLE DE LA OFERTA" header bar
header_bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.5), Cm(3.2), Cm(28.7), Cm(0.8))
header_bar.fill.solid()
header_bar.fill.fore_color.rgb = GREEN
header_bar.line.fill.background()

add_text_box(slide2, Cm(0.7), Cm(3.3), Cm(28.3), Cm(0.6), "DETALLE DE LA OFERTA", 14, True, WHITE, PP_ALIGN.LEFT)

# Table columns headers
col_headers = ["", "TIPO", "SERVICIO/EQUIPO", "COMPROMISO", "CANTIDAD", "PRECIO", "DTO", "CUOTA ALTA", "DTO C. ALTA", "TOTAL"]
col_widths = [2.5, 2.5, 7.0, 3.0, 2.5, 2.5, 2.0, 2.5, 2.5, 2.7]
col_starts = [0.5]
for w in col_widths[:-1]:
    col_starts.append(col_starts[-1] + w)

header_y = Cm(4.0)
for i, (header, x, w) in enumerate(zip(col_headers, col_starts, col_widths)):
    add_text_box(slide2, Cm(x), header_y, Cm(w - 0.2), Cm(0.5), header, 8, True, BLACK, PP_ALIGN.CENTER if i > 0 else PP_ALIGN.LEFT)

# Table rows data (Jose Vicente defaults)
table_data = [
    ("PAQUETES", "", "B2ONE B8 PREMIUM", "36 meses", "1", "149,00€/mes", "", "", "", "149,00 €"),
    ("", "", "LYNKS FTTO 1GB + BACKUP 4G", "36 meses", "1", "", "", "", "", "0,00 €"),
    ("", "", "CANAL VOZ (TRUNK)", "36 meses", "8", "", "", "", "", "0,00 €"),
    ("", "", "EXT IP (IVR,GRABACIÓN,SOFTPHONE)", "36 meses", "8", "", "", "", "", "0,00 €"),
    ("", "", "DDI (NACIONAL)", "36 meses", "8", "", "", "", "", "0,00 €"),
    ("", "", "MOVIL B2 ILIMITADA PACK", "36 meses", "2", "", "", "", "", "0,00 €"),
    ("LYNKS MOBILE", "TARIFAS", "MOVIL B2 ILIMITADA PACK", "12 meses", "4", "9,45€/mes", "", "", "", "37,80 €"),
    ("LYNKS IP", "EQUIPOS Y TERMINALES", "GIGASET BÁSICO P710", "36 meses", "8", "3,00€/mes", "", "", "", "24,00 €"),
]

# Total row
total_row = ["", "", "", "", "", "", "", "", "TOTAL PRIMER PAGO", "210,80 €"]

for row_idx, row_data in enumerate(table_data):
    y = Cm(4.6 + row_idx * 0.65)
    bg_color = TABLE_ROW_BG if row_idx % 2 == 0 else TABLE_ALT_BG
    
    # Row background
    row_bg = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.5), y, Cm(28.7), Cm(0.6))
    row_bg.fill.solid()
    row_bg.fill.fore_color.rgb = bg_color
    row_bg.line.fill.background()
    
    for col_idx, (cell, x, w) in enumerate(zip(row_data, col_starts, col_widths)):
        add_text_box(slide2, Cm(x + 0.1), y + Cm(0.1), Cm(w - 0.2), Cm(0.4), cell, 8, False, BLACK, PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT)

# Total row background
total_y = Cm(4.6 + len(table_data) * 0.65 + 0.3)
total_bg = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.5), total_y, Cm(28.7), Cm(0.8))
total_bg.fill.solid()
total_bg.fill.fore_color.rgb = TOTAL_BG
total_bg.line.fill.background()

# Total labels and values
add_text_box(slide2, Cm(12), total_y + Cm(0.15), Cm(8), Cm(0.5), "TOTAL PRIMER PAGO", 12, True, WHITE, PP_ALIGN.RIGHT)
add_text_box(slide2, Cm(20.5), total_y + Cm(0.15), Cm(3), Cm(0.5), "210,80 €", 14, True, WHITE, PP_ALIGN.RIGHT)
add_text_box(slide2, Cm(23.5), total_y + Cm(0.15), Cm(5), Cm(0.5), "RESTO MEN...", 10, True, WHITE, PP_ALIGN.RIGHT)

# Footer note
add_text_box(slide2, Cm(0.5), Cm(19.5), Cm(20), Cm(0.4), "* Los precios no incluyen IVA / IGIC", 8, False, MID_GRAY, PP_ALIGN.LEFT)

# ============================
# SLIDE 3: CLOUD / BACKUP / CIBERSEGURIDAD
# ============================
slide3 = prs.slides.add_slide(blank_layout)

# Dark background
bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(29.7), Cm(21.0))
bg3.fill.solid()
bg3.fill.fore_color.rgb = DARK_BG
bg3.line.fill.background()

# Top green bar
top_bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(29.7), Cm(0.8))
top_bar.fill.solid()
top_bar.fill.fore_color.rgb = GREEN
top_bar.line.fill.background()

# Bottom green bar
bot_bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(20.2), Cm(29.7), Cm(0.8))
bot_bar.fill.solid()
bot_bar.fill.fore_color.rgb = GREEN
bot_bar.line.fill.background()

# Title
add_text_box(slide3, Cm(0), Cm(1.5), Cm(29.7), Cm(1.2), "CLOUD · BACKUP · CIBERSEGURIDAD", 32, True, WHITE, PP_ALIGN.CENTER)

# Subtitle
add_text_box(slide3, Cm(0), Cm(2.8), Cm(29.7), Cm(0.5), "Infraestructura en CPD Tier-IV Madrid · Distribuidor oficial Datos101", 11, False, LIGHT_GRAY, PP_ALIGN.CENTER)

# Certification badges
badges = ["ISO 27001", "ISO 20000", "ENS", "RGPD", "NIS2-ready", "SLA 99,99%"]
for i, badge in enumerate(badges):
    x = Cm(2 + i * 4.5)
    badge_shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Cm(3.8), Cm(4.2), Cm(0.7))
    badge_shape.fill.solid()
    badge_shape.fill.fore_color.rgb = GREEN
    badge_shape.line.fill.background()
    add_text_box(slide3, x, Cm(3.85), Cm(4.2), Cm(0.6), badge, 10, True, WHITE, PP_ALIGN.CENTER)

# 4 service boxes
services = [
    ("CLOUD & IaaS", "Cloud101 · VPS · Infraestructura", ["VPS y máquinas virtuales on-demand", "Escalado flexible · pago por uso", "Soberanía del dato en España / UE"]),
    ("BACKUP", "Backup as a Service · Backup365", ["Copia cifrada de servidores, M365 y NAS", "Doble destino (local + cloud)", "Restauración granular en minutos"]),
    ("DISASTER RECOVERY", "DRaaS · Plan de contingencia", ["Replicación continua RTO/RPO bajos", "Simulacros periódicos de failover", "Réplicas inmutables anti-ransomware"]),
    ("CIBERSEGURIDAD", "Ciber101 · SOC gestionado", ["SOC 24×7 · detección y respuesta", "EDR/XDR · endpoint y servidor", "Auditorías técnicas · cumplimiento NIS2"]),
]

for i, (title, subtitle, bullets) in enumerate(services):
    col = i % 2
    row = i // 2
    x = Cm(1.5 + col * 14)
    y = Cm(5.2 + row * 5.5)
    
    # Box with green left border
    box = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Cm(13), Cm(4.8))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    box.line.color.rgb = RGBColor(0x3A, 0x3A, 0x3A)
    box.line.width = Pt(1)
    
    # Green left border
    left_border = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Cm(0.3), Cm(4.8))
    left_border.fill.solid()
    left_border.fill.fore_color.rgb = GREEN
    left_border.line.fill.background()
    
    add_text_box(slide3, x + Cm(0.8), y + Cm(0.3), Cm(11.5), Cm(0.6), title, 14, True, WHITE, PP_ALIGN.LEFT)
    add_text_box(slide3, x + Cm(0.8), y + Cm(0.9), Cm(11.5), Cm(0.4), subtitle, 9, False, GREEN, PP_ALIGN.LEFT)
    
    for j, bullet in enumerate(bullets):
        add_text_box(slide3, x + Cm(1.2), y + Cm(1.6 + j * 0.7), Cm(10.5), Cm(0.5), f"■  {bullet}", 9, False, LIGHT_GRAY, PP_ALIGN.LEFT)

# Footer bar with contact
footer_bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.5), Cm(18.5), Cm(26.7), Cm(1.0))
footer_bar.fill.solid()
footer_bar.fill.fore_color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
footer_bar.line.color.rgb = GREEN
footer_bar.line.width = Pt(2)

add_text_box(slide3, Cm(2), Cm(18.7), Cm(25), Cm(0.6), "HABLEMOS  ·  911 08 98 77  ·  info@lynks-tic.com  ·  lynks-tic.com", 11, True, WHITE, PP_ALIGN.LEFT)

# Page number
add_text_box(slide3, Cm(27), Cm(19.5), Cm(2), Cm(0.4), "05 / 05", 8, False, MID_GRAY, PP_ALIGN.RIGHT)

# Save
output_path = "/Users/javiermartinamador/Documents/Agencia Marketing/templates/lynks-tic_pro_base.pptx"
prs.save(output_path)
print(f"Template saved: {output_path}")
print(f"Size: {len(prs.slides)} slides")
