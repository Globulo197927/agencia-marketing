import shutil
from pathlib import Path
from datetime import datetime
from pptx import Presentation
from pptx.util import Cm
from lxml import etree

from src.config_loader import load_brand_config, BASE_DIR


def _replace_text_in_shape(shape, find_text: str, replace_text: str) -> bool:
    if not find_text or find_text == replace_text:
        return False
    if not shape.has_text_frame:
        return False

    for paragraph in shape.text_frame.paragraphs:
        full_text = "".join(run.text for run in paragraph.runs)
        if find_text in full_text:
            new_text = full_text.replace(find_text, replace_text)
            if paragraph.runs:
                paragraph.runs[0].text = new_text
                for run in paragraph.runs[1:]:
                    run.text = ""
                return True
    return False


def _replace_in_all_shapes(prs: Presentation, find_text: str, replace_text: str) -> bool:
    if find_text == replace_text:
        return False
    found = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if _replace_text_in_shape(shape, find_text, replace_text):
                found = True
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            full_text = "".join(run.text for run in paragraph.runs)
                            if find_text in full_text:
                                new_text = full_text.replace(find_text, replace_text)
                                if paragraph.runs:
                                    paragraph.runs[0].text = new_text
                                    for run in paragraph.runs[1:]:
                                        run.text = ""
                                    found = True
    return found


def _replace_in_slide(prs: Presentation, slide_idx: int, find_text: str, replace_text: str) -> bool:
    if find_text == replace_text:
        return False
    slide = prs.slides[slide_idx]
    found = False
    for shape in slide.shapes:
        if _replace_text_in_shape(shape, find_text, replace_text):
            found = True
    return found


def _replace_table_cell(prs: Presentation, slide_idx: int, find_text: str, replace_text: str):
    if find_text == replace_text:
        return
    slide = prs.slides[slide_idx]
    for shape in slide.shapes:
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        full_text = "".join(run.text for run in paragraph.runs)
                        if find_text in full_text:
                            new_text = full_text.replace(find_text, replace_text)
                            if paragraph.runs:
                                paragraph.runs[0].text = new_text
                                for run in paragraph.runs[1:]:
                                    run.text = ""


def _remove_slide(prs: Presentation, slide_idx: int):
    rId = prs.slides._sldIdLst[slide_idx].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    sldId = prs.slides._sldIdLst[slide_idx]
    prs.slides._sldIdLst.remove(sldId)


def _add_slides_from_template(prs: Presentation, template_path: str, start_slide: int, end_slide: int):
    src_prs = Presentation(template_path)
    for i in range(start_slide - 1, min(end_slide, len(src_prs.slides))):
        src_slide = src_prs.slides[i]
        slide_layout = prs.slide_layouts[0]
        new_slide = prs.slides.add_slide(slide_layout)
        for shape in src_slide.shapes:
            el = shape._element
            new_slide.shapes._spTree.append(el)


TERMINAL_IMAGES = {
    "vodafone": {
        "slide": 7,
        "terminals": [
            {"image": "assets/adoc_neo_8000.png", "x": "5.2cm", "y": "2.7cm", "w": "2.5cm", "h": "2.5cm"},
            {"image": "assets/adoc_neo_3850w.png", "x": "13.6cm", "y": "2.7cm", "w": "2.5cm", "h": "2.5cm"},
            {"image": "assets/cocomm_dt250.png", "x": "21.8cm", "y": "2.7cm", "w": "2.5cm", "h": "2.5cm"},
        ],
    },
    "lynks-tic": {
        "slide": 12,
        "terminals": [
            {"image": "assets/yealink_w73p.png", "replace_name": "Image 0", "x": "1.83cm", "y": "3.05cm", "w": "6.1cm", "h": "10.16cm"},
        ],
    },
}


def _parse_cm(value: str) -> int:
    return Cm(float(value.replace("cm", "")))


def _add_terminal_images(prs: Presentation, client_data: dict, brand: str, base_dir: Path):
    config = TERMINAL_IMAGES.get(brand)
    if not config:
        return
    slide = prs.slides[config["slide"] - 1]

    if brand == "vodafone":
        for term_conf in config["terminals"]:
            ipath = base_dir / term_conf["image"]
            if ipath.exists():
                slide.shapes.add_picture(
                    str(ipath),
                    _parse_cm(term_conf["x"]),
                    _parse_cm(term_conf["y"]),
                    _parse_cm(term_conf["w"]),
                    _parse_cm(term_conf["h"]),
                )
    elif brand == "lynks-tic":
        for term_conf in config["terminals"]:
            ipath = base_dir / term_conf["image"]
            if not ipath.exists():
                continue
            replace_name = term_conf.get("replace_name")
            if replace_name:
                for shape in list(slide.shapes):
                    try:
                        if hasattr(shape, 'image') and shape.name == replace_name:
                            sp = shape._element
                            sp.getparent().remove(sp)
                            break
                    except Exception:
                        pass
            slide.shapes.add_picture(
                str(ipath),
                _parse_cm(term_conf["x"]),
                _parse_cm(term_conf["y"]),
                _parse_cm(term_conf["w"]),
                _parse_cm(term_conf["h"]),
            )


def _get_client_name(data: dict, brand: str) -> str:
    if brand == "vodafone":
        c = data.get("client", {})
        return f"{c.get('name_line1', 'Cliente')}_{c.get('name_line2', '')}".strip("_")
    c = data.get("client", {})
    return f"{c.get('name_short', 'Cliente')}_{c.get('name_suffix', '')}".strip("_")


def generate_pptx(brand: str, client_data: dict, output_name: str = None) -> Path:
    config = load_brand_config(brand)
    template_path = BASE_DIR / config["brand"]["template"]
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    brand_dir = brand.replace("-", "_")
    if brand_dir not in ("vodafone", "lynks_tic", "lynks_tic_b2mobile"):
        brand_dir = "lynks-tic"
    output_dir = BASE_DIR / "output" / brand_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    if output_name is None:
        client_name = _get_client_name(client_data, brand)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_name = f"{brand}_{client_name}_{timestamp}.pptx"

    if not output_name.endswith(".pptx"):
        output_name += ".pptx"

    output_path = output_dir / output_name
    shutil.copy2(template_path, output_path)

    prs = Presentation(str(output_path.resolve()))

    if brand == "vodafone":
        _apply_vodafone_data(prs, client_data)
    elif brand == "lynks-tic-b2mobile":
        _apply_b2mobile_data(prs, client_data)
    elif brand == "lynks-tic-pro":
        _apply_pro_data(prs, client_data)
    else:
        _apply_lynks_data(prs, client_data)

    _add_terminal_images(prs, client_data, brand, BASE_DIR)

    _handle_optional_slides(prs, output_path, client_data, brand)

    prs.save(str(output_path.resolve()))

    return output_path


def _apply_vodafone_data(prs: Presentation, data: dict):
    c = data.get("client", {})
    p = data.get("proposal", {})
    prods = data.get("products", [])
    eco = data.get("economic", {})
    terms = data.get("terminals", [])

    full_name = f"{c.get('name_line1', '')} {c.get('name_line2', '')}".strip()
    lines = c.get("lines", 25)
    ctype = c.get("type", "Cliente Nuevo")
    date = p.get("date", "Marzo 2026")
    net = p.get("network_type", "Red 5G Vodafone")

    # Slide 1
    _replace_in_slide(prs, 0, "TECNOLOGÍA", c.get("name_line1", "TECNOLOGÍA"))
    _replace_in_slide(prs, 0, "TÉRMICA SA", c.get("name_line2", "TÉRMICA SA"))
    _replace_in_slide(prs, 0, "NIF A28592772", f"NIF {c.get('nif', 'A28592772')}")
    _replace_in_slide(prs, 0, "● Red 5G Vodafone", f"● {net}")
    _replace_in_slide(prs, 0, "Abril 2026", date)

    bullets = c.get("feature_bullets", "")
    if bullets:
        _replace_in_slide(prs, 0, "▸  25 líneas móviles · Plan RED Infinity PRO", bullets.split("\n")[0] if bullets else "")
    elif prods:
        p1 = prods[0]
        new_line1 = f"▸  {lines} líneas móviles · {p1.get('name_line1', '')} {p1.get('name_line2', '')}"
        _replace_in_slide(prs, 0, "▸  25 líneas móviles · Plan RED Infinity PRO", new_line1)

    # Slide 2
    _replace_in_slide(prs, 1, "TECNOLOGÍA TÉRMICA SA", full_name)

    # Slide 5
    summary = f"Servicios base para {full_name} · {lines} líneas · {ctype} · Compromiso {p.get('duration_months', '36 meses')}"
    _replace_in_slide(prs, 4, "Servicios base para TECNOLOGÍA TÉRMICA SA · 25 líneas · Cliente Nuevo · Compromiso 36 meses", summary)

    if len(prods) >= 1:
        p1 = prods[0]
        _replace_in_slide(prs, 4, "BUNDLE RED\nINFINITY PRO", f"{p1.get('name_line1', 'BUNDLE RED')}\n{p1.get('name_line2', 'INFINITY PRO')}")
        _replace_in_slide(prs, 4, "73,44€/mes", p1.get("price_monthly", "73,44€/mes"))
        _replace_in_slide(prs, 4, "antes: 112,98€/mes", f"antes: {p1.get('price_before', '112,98€/mes')}")
        _replace_in_slide(prs, 4, "35% de descuento", p1.get("discount_text", "35% de descuento"))
        feats = p1.get("features", [])
        default_feats_1 = ["Datos ilimitados Full Speed", "5G en toda la red Vodafone", "Roaming incluido en UE"]
        for i in range(min(3, len(feats))):
            _replace_in_slide(prs, 4, default_feats_1[i], feats[i])

    if len(prods) >= 2:
        p2 = prods[1]
        _replace_in_slide(prs, 4, "FIBRA\n1 GBPS", f"{p2.get('name_line1', 'FIBRA')}\n{p2.get('name_line2', '1 GBPS')}")
        _replace_in_slide(prs, 4, "0€/mes", p2.get("price_monthly", "0€/mes"), )
        _replace_in_slide(prs, 4, "antes: 34,85€/mes", f"antes: {p2.get('price_before', '34,85€/mes')}")
        _replace_in_slide(prs, 4, "100% de descuento", p2.get("discount_text", "100% de descuento"))
        feats2 = p2.get("features", [])
        default_feats_2 = ["Conexión simétrica 1 Gbps", "Incluida sin coste adicional", "Instalación y mantenimiento"]
        for i in range(min(3, len(feats2))):
            _replace_in_slide(prs, 4, default_feats_2[i], feats2[i])

    if len(prods) >= 3:
        p3 = prods[2]
        _replace_in_slide(prs, 4, "CENTRALITA\nPLUS", f"{p3.get('name_line1', 'CENTRALITA')}\n{p3.get('name_line2', 'PLUS')}")
        _replace_in_slide(prs, 4, "antes: 30,75€/mes", f"antes: {p3.get('price_before', '30,75€/mes')}")
        feats3 = p3.get("features", [])
        default_feats_3 = ["Gestión avanzada de llamadas", "Incluida sin coste adicional", "Multi-sede y extensiones"]
        for i in range(min(3, len(feats3))):
            _replace_in_slide(prs, 4, default_feats_3[i], feats3[i])

    _replace_in_slide(prs, 4, "25 líneas móviles", f"{lines} líneas móviles")
    _replace_in_slide(prs, 4, "Red 5G Vodafone", net)
    _replace_in_slide(prs, 4, "Cliente Nuevo", ctype)
    _replace_in_slide(prs, 4, "1.800€", eco.get("apoyo_economico", "1.800€"))

    # Slide 6
    if eco:
        _replace_in_slide(prs, 5, "305,79 €/mes", eco.get("total_monthly", "305,79 €/mes"))
        if eco.get("summary_line"):
            _replace_in_all_shapes(prs, "25 líneas · Terminales a 0€ · Todos con 35% dto · Compromiso 36 meses", eco["summary_line"])
        if eco.get("subsidy_detail_line"):
            _replace_in_slide(prs, 5, "Sub. Vodafone: -160,82€", eco["subsidy_detail_line"].split("·")[0].strip())

    # Slide 6 - Table
    table_data = data.get("distribution_table", [])
    if table_data:
        default_table = [
            {"terminal": "Adoc NEO 8000 4G Deals", "plan": "RED Infinity PRO Ilim. Bundle", "precio_cesion": "187,00€", "sin_dto": "0,00€", "dto": "35%", "con_dto": "0,00€", "uds": "1", "total": "0,00€"},
            {"terminal": "Adoc NEO 3850W 4G Deals", "plan": "RED Infinity PRO Ilim. Full Speed", "precio_cesion": "365,00€", "sin_dto": "17,95€", "dto": "35%", "con_dto": "11,67€", "uds": "5", "total": "58,34€"},
            {"terminal": "Adoc NEO 3850W 4G Deals", "plan": "RED Infinity PRO Ilim. Full Speed", "precio_cesion": "365,00€", "sin_dto": "17,95€", "dto": "35%", "con_dto": "11,67€", "uds": "5", "total": "58,34€"},
            {"terminal": "—", "plan": "RED Infinity PRO 5Gb", "precio_cesion": "0,00€", "sin_dto": "13,73€", "dto": "35%", "con_dto": "8,92€", "uds": "10", "total": "89,24€"},
            {"terminal": "—", "plan": "RED Infinity PRO 5Gb", "precio_cesion": "0,00€", "sin_dto": "13,73€", "dto": "35%", "con_dto": "8,92€", "uds": "1", "total": "8,92€"},
            {"terminal": "Cocomm DT250 4G Solo Voz", "plan": "RED Infinity PRO Solo Voz", "precio_cesion": "198,00€", "sin_dto": "8,98€", "dto": "35%", "con_dto": "5,84€", "uds": "3", "total": "17,51€"},
        ]
        for row_idx, row_data in enumerate(table_data):
            if row_idx >= 6:
                break
            default_row = default_table[row_idx] if row_idx < len(default_table) else {}
            for key in ["terminal", "plan", "precio_cesion", "sin_dto", "dto", "con_dto", "uds", "total"]:
                default_value = default_row.get(key, "")
                new_value = row_data.get(key, default_value)
                if default_value and new_value:
                    _replace_in_all_shapes(prs, default_value, new_value)

    # Slide 7 - Terminals
    if terms:
        for i, t in enumerate(terms[:3]):
            defaults = ["Adoc NEO 8000 4G", "Adoc NEO 3850W 4G", "Cocomm DT250 4G Solo Voz"]
            _replace_in_slide(prs, 6, defaults[i], t.get("model", defaults[i]))
            if t.get("price"):
                _replace_in_slide(prs, 6, "PRECIO CESIÓN: 0,00 €", f"PRECIO CESIÓN: {t['price']}")

    # Slide 8
    _replace_in_slide(prs, 7, "TECNOLOGÍA TÉRMICA SA", full_name)


def _apply_lynks_data(prs: Presentation, data: dict):
    c = data.get("client", {})
    p = data.get("proposal", {})
    opts = data.get("options", [])

    short = c.get("name_short", "DIAVERUM")
    suffix = c.get("name_suffix", "SERVICIOS RENALES SL")
    full = f"{short} {suffix}".strip()
    sedes = c.get("sedes", 55)
    nif = c.get("nif", "B84991736")
    date = p.get("date", "Marzo 2026")
    ext = p.get("extensiones", 165)

    # Slide 1
    _replace_in_slide(prs, 0, "DIAVERUM", short)
    _replace_in_slide(prs, 0, "SERVICIOS RENALES SL", suffix)
    _replace_in_slide(prs, 0, "B84991736  ·  55 sedes  ·  Conectividad FTTO + Voz IP  ·  Marzo 2026",
                       f"{nif}  ·  {sedes} sedes  ·  {c.get('service_type', 'Conectividad FTTO + Voz IP')}  ·  {date}")

    # Slides 2, 4
    _replace_in_slide(prs, 1, "Diaverum", short)
    _replace_in_slide(prs, 3, "Diaverum Servicios Renales SL", full)

    # Slide 12
    _replace_in_slide(prs, 11, "165 unidades · 8,50€/ud/mes · Tarifa plana ilimitada a fijos y móviles nacionales",
                       f"{ext} unidades · {p.get('terminal_unit_price', '8,50€/ud/mes')} · Tarifa plana ilimitada a fijos y móviles nacionales")

    # Slide 13
    _replace_in_slide(prs, 12, "las 55 sedes de Diaverum Servicios Renales SL", f"las {sedes} sedes de {full}")
    _replace_in_slide(prs, 12, "165", str(ext))
    _replace_in_slide(prs, 12, "55", str(sedes))

    # Slide 14 - Option A
    if len(opts) >= 1:
        a = opts[0]
        _replace_in_slide(prs, 13, "OPCIÓN A — FTTO 1 GBPS · LA MÁS POTENTE",
                           a.get("title", "OPCIÓN A — FTTO 1 GBPS · LA MÁS POTENTE"))
        _replace_in_slide(prs, 13, "3.795 €/mes", a.get("total_price", "3.795 €/mes"))
        _replace_in_slide(prs, 13, "55 sedes · 1Gbps/1Gbps simétrico · Compromiso anual · SLA 99,9%",
                           f"{sedes} sedes · {a.get('speed', '1Gbps/1Gbps simétrico')} · Compromiso anual · SLA 99,9%")
        _replace_in_slide(prs, 13, "Descarga 1.000 Mbps · Subida 1.000 Mbps · Latencia <10ms",
                           f"Descarga {a.get('speed_download', '1.000 Mbps')} · Subida {a.get('speed_upload', '1.000 Mbps')} · Latencia <10ms")

    # Slide 15 - Option B
    if len(opts) >= 2:
        b = opts[1]
        _replace_in_slide(prs, 14, "OPCIÓN B — FTTO 600 MBPS · LA MÁS EQUILIBRADA",
                           b.get("title", "OPCIÓN B — FTTO 600 MBPS · LA MÁS EQUILIBRADA"))
        _replace_in_slide(prs, 14, "3.575 €/mes", b.get("total_price", "3.575 €/mes"))
        _replace_in_slide(prs, 14, "55 sedes · 600Mbps/600Mbps simétrico · Compromiso anual · SLA 99,9%",
                           f"{sedes} sedes · {b.get('speed', '600Mbps/600Mbps simétrico')} · Compromiso anual · SLA 99,9%")

    # Slide 16 - Comparison
    if len(opts) >= 1:
        a = opts[0]
        _replace_in_slide(prs, 15, "FTTO 1 Gbps simétrico", a.get("speed_label", "FTTO 1 Gbps simétrico"))
        _replace_in_slide(prs, 15, "3.795 €/mes", a.get("total_price", "3.795 €/mes"))
        _replace_in_slide(prs, 15, "55 × FTTO 1 Gbps — 2.310€", f"{sedes} × FTTO 1 Gbps — {a.get('ftto_total', '2.310€')}")

    if len(opts) >= 2:
        b = opts[1]
        _replace_in_slide(prs, 15, "FTTO 600 Mbps simétrico", b.get("speed_label", "FTTO 600 Mbps simétrico"))
        _replace_in_slide(prs, 15, "3.575 €/mes", b.get("total_price", "3.575 €/mes"))
        _replace_in_slide(prs, 15, "55 × FTTO 600 Mbps — 2.090€", f"{sedes} × FTTO 600 Mbps — {b.get('ftto_total', '2.090€')}")

    _replace_in_slide(prs, 15, "55 × DDI — 82,50€", f"{sedes} × DDI — {sedes * 1.5:.2f}€")
    _replace_in_slide(prs, 15, "165 × Yealink W37P — 1.402,50€", f"{ext} × Yealink W37P — {ext * 8.5:.2f}€")
    _replace_in_slide(prs, 15, "↓ Ahorro de 2.640€/año respecto a Opción A",
                       data.get("savings_line", "↓ Ahorro de 2.640€/año respecto a Opción A"))

    # Slide 17
    _replace_in_slide(prs, 16, "DIAVERUM SERVICIOS RENALES SL", full)


def _apply_b2mobile_data(prs: Presentation, data: dict):
    c = data.get("client", {})
    p = data.get("proposal", {})
    comm = data.get("commercial", {})
    conn = data.get("connectivity", [])
    cyber = data.get("cybersecurity", {})

    full_name = f"{c.get('name_short', '')} {c.get('name_suffix', '')}".strip()
    lines = c.get("lines", 730)
    nif = c.get("nif", "")
    date = p.get("date", "Junio 2026")
    duration = p.get("duration_months", 36)
    pack_name = p.get("pack_name", "B2 Pack Ilimitada")
    price_per_line = p.get("price_per_line", "7,50€/línea/mes")
    total_monthly = p.get("total_monthly", "5.475 €/mes")
    total_before = p.get("total_before", "6.898,50€/mes")
    discount_text = p.get("discount_text", "21% de descuento")
    catalog_price = p.get("catalog_price", "9,45€")
    bag_terminals = p.get("bag_terminals", "8.000€")

    # Slide 1 — Cover
    _replace_in_slide(prs, 0, "HIJAS DE MARÍA AUXILIADORA", c.get("name_short", "HIJAS DE MARÍA AUXILIADORA"))
    _replace_in_slide(prs, 0, "INSTITUTO SALESIANAS CIR", c.get("name_suffix", ""))
    _replace_in_slide(prs, 0, "NIF R0800057B", f"NIF {nif}")
    _replace_in_slide(prs, 0, "730 líneas móviles", f"{lines} líneas móviles")
    _replace_in_slide(prs, 0, "B2 Pack Ilimitada", pack_name)
    _replace_in_slide(prs, 0, "Junio 2026", date)
    _replace_in_slide(prs, 0, "8.000€", bag_terminals)

    # Slide 3 — B2Mobile detail
    _replace_in_slide(prs, 2, "730", str(lines))
    _replace_in_slide(prs, 2, "7,50€/línea/mes", price_per_line)
    _replace_in_slide(prs, 2, "9,45€", catalog_price)
    _replace_in_slide(prs, 2, "21% de descuento", discount_text)
    _replace_in_slide(prs, 2, f"para {lines} líneas", f"para {lines} líneas")

    # Slide 4 — What's included
    _replace_in_all_shapes(prs, "HIJAS DE MARÍA AUXILIADORA", c.get("name_short", "HIJAS DE MARÍA AUXILIADORA"))
    _replace_in_all_shapes(prs, "INSTITUTO SALESIANAS CIR", c.get("name_suffix", ""))
    _replace_in_slide(prs, 3, "5.475€/mes", total_monthly.replace(" ", ""))
    _replace_in_slide(prs, 3, "antes: 6.898,50€/mes", f"antes: {total_before}")
    _replace_in_slide(prs, 3, "21% de descuento", discount_text)
    _replace_in_slide(prs, 3, "730 líneas móviles", f"{lines} líneas móviles")
    _replace_in_slide(prs, 3, "7,50€/línea/mes", price_per_line)
    _replace_in_slide(prs, 3, "8.000€", bag_terminals)

    # Slide 5 — Economic detail
    _replace_in_slide(prs, 4, "5.475 €/mes", total_monthly)
    _replace_in_slide(prs, 4, "730 líneas · 7,50€/línea", f"{lines} líneas · {price_per_line}")
    _replace_in_all_shapes(prs, "HIJAS DE MARÍA AUXILIADORA · INSTITUTO SALESIANAS CIR", full_name)
    _replace_in_slide(prs, 4, "R0800057B", nif)
    _replace_in_slide(prs, 4, "36 meses", f"{duration} meses")
    _replace_in_slide(prs, 4, "5.475,00€", total_monthly.replace(" ", "").replace("€", ",00€") if "5475" in total_monthly else "5.475,00€")
    _replace_in_slide(prs, 4, "7,50€", price_per_line.replace("/línea/mes", ""))
    _replace_in_slide(prs, 4, "9,45€", catalog_price)
    _replace_in_slide(prs, 4, "730 uds.", f"{lines} uds.")
    _replace_in_slide(prs, 4, "8.000,00€", bag_terminals)

    # Slide 6 — Connectivity options
    if conn:
        ftto_names = ["FTTO 1Gb", "FTTO 600Mb", "FTTO 300Mb"]
        ftto_prices = ["69€/mes", "49€/mes", "34€/mes"]
        for i, ftto in enumerate(conn[:3]):
            _replace_in_slide(prs, 5, ftto_names[i], ftto.get("name", ftto_names[i]))
            _replace_in_slide(prs, 5, ftto_prices[i], ftto.get("price", ftto_prices[i]))

    # Slide 7 — Cybersecurity
    if cyber:
        _replace_in_slide(prs, 6, "70€/mes", cyber.get("pack_price", "70€/mes"))
        _replace_in_slide(prs, 6, "35€/mes", cyber.get("extra_price", "35€/mes"))

    # Slide 8 — Next steps
    _replace_in_all_shapes(prs, "HIJAS DE MARÍA AUXILIADORA", c.get("name_short", "HIJAS DE MARÍA AUXILIADORA"))
    _replace_in_slide(prs, 7, "730 líneas móviles", f"{lines} líneas móviles")

    if comm:
        comm_name = comm.get("name", "Javier Martín Amador")
        comm_phone = comm.get("phone", "664 25 89 77")
        comm_email = comm.get("email", "javier.martin@lynks-tic.com")
        _replace_in_slide(prs, 7, "Javier Martín Amador", comm_name)
        _replace_in_slide(prs, 7, "664 25 89 77", comm_phone)
        _replace_in_slide(prs, 7, "javier.martin@lynks-tic.com", comm_email)

    # Footer on all slides
    for i in range(len(prs.slides)):
        _replace_in_slide(prs, i, "Junio 2026", date)
        if comm:
            _replace_in_slide(prs, i, "Javier Martín", comm.get("name", "Javier Martín").split()[0])
            _replace_in_slide(prs, i, "664 25 89 77", comm.get("phone", "664 25 89 77"))


def _apply_pro_data(prs: Presentation, data: dict):
    c = data.get("client", {})
    comm = data.get("commercial", {})
    svcs = data.get("services", [])
    total = data.get("total", {})

    client_name = c.get("name", "JOSE VICENTE SL")
    nif = c.get("nif", "B28927994")
    full_client = f"{client_name} ({nif})"

    comm_name = comm.get("name", "JAVIER MARTIN AMADOR")
    comm_phone = comm.get("phone", "664258977")
    comm_email = comm.get("email", "JAVIER.MARTIN@LYNKS-TIC.COM")

    # Slide 1 — Cover
    _replace_in_slide(prs, 0, "JOSE VICENTE SL (B28927994)", full_client)

    # Slide 2 — Detail table
    _replace_in_slide(prs, 1, "JOSE VICENTE SL (B28927994)", full_client)
    _replace_in_slide(prs, 1, "JAVIER MARTIN AMADOR", comm_name)
    _replace_in_slide(prs, 1, "664258977", comm_phone)
    _replace_in_slide(prs, 1, "JAVIER.MARTIN@LYNKS-TIC.COM", comm_email)

    # Replace table rows
    default_services = [
        ("PAQUETES", "", "B2ONE B8 PREMIUM", "36 meses", "1", "149,00€/mes", "", "", "", "149,00 €"),
        ("", "", "LYNKS FTTO 1GB + BACKUP 4G", "36 meses", "1", "", "", "", "", "0,00 €"),
        ("", "", "CANAL VOZ (TRUNK)", "36 meses", "8", "", "", "", "", "0,00 €"),
        ("", "", "EXT IP (IVR,GRABACIÓN,SOFTPHONE)", "36 meses", "8", "", "", "", "", "0,00 €"),
        ("", "", "DDI (NACIONAL)", "36 meses", "8", "", "", "", "", "0,00 €"),
        ("", "", "MOVIL B2 ILIMITADA PACK", "36 meses", "2", "", "", "", "", "0,00 €"),
        ("LYNKS MOBILE", "TARIFAS", "MOVIL B2 ILIMITADA PACK", "12 meses", "4", "9,45€/mes", "", "", "", "37,80 €"),
        ("LYNKS IP", "EQUIPOS Y TERMINALES", "GIGASET BÁSICO P710", "36 meses", "8", "3,00€/mes", "", "", "", "24,00 €"),
    ]

    services_to_use = svcs if svcs else []
    while len(services_to_use) < len(default_services):
        idx = len(services_to_use)
        services_to_use.append({
            "category": default_services[idx][0],
            "type": default_services[idx][1],
            "name": default_services[idx][2],
            "commitment": default_services[idx][3],
            "quantity": default_services[idx][4],
            "price": default_services[idx][5],
            "discount": default_services[idx][6],
            "cuota_alta": default_services[idx][7],
            "dto_cuota": default_services[idx][8],
            "total": default_services[idx][9],
        })

    for i, svc in enumerate(services_to_use[:len(default_services)]):
        default_row = default_services[i]
        if svc.get("name") and svc["name"] != default_row[2]:
            _replace_in_slide(prs, 1, default_row[2], svc["name"])
        if svc.get("commitment") and svc["commitment"] != default_row[3]:
            _replace_in_slide(prs, 1, default_row[3], svc["commitment"])
        if svc.get("quantity") and svc["quantity"] != default_row[4]:
            _replace_in_slide(prs, 1, default_row[4], svc["quantity"])
        if svc.get("price") and svc["price"] != default_row[5]:
            _replace_in_slide(prs, 1, default_row[5], svc["price"])
        if svc.get("total") and svc["total"] != default_row[9]:
            _replace_in_slide(prs, 1, default_row[9], svc["total"])

    # Replace totals
    primer_pago = total.get("primer_pago", "210,80 €")
    mensualidades = total.get("mensualidades", "210,80 €/mes")
    _replace_in_slide(prs, 1, "210,80 €", primer_pago)
    _replace_in_slide(prs, 1, "RESTO MEN...", f"RESTO MENSUALIDADES {mensualidades}")


def _handle_optional_slides(prs: Presentation, output_path: Path, client_data: dict, brand: str):
    if brand != "vodafone":
        return

    proposal = client_data.get("proposal", client_data)
    include_holidaysim = proposal.get("include_holidaysim", True)
    include_lynks_tic = proposal.get("include_lynks_tic", True)

    if not include_holidaysim and not include_lynks_tic:
        _remove_slide(prs, 2)
        _remove_slide(prs, 2)
    elif not include_holidaysim:
        _remove_slide(prs, 2)
    elif not include_lynks_tic:
        _remove_slide(prs, 3)